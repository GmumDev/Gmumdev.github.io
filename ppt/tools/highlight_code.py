#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""포트폴리오 pptx 코드 박스에 구문 강조를 입힌다.

규칙과 팔레트는 ppt/CODE_THEME.md 에 적혀 있다. 이 파일은 그 문서의 구현이다.

사용:
    python ppt/tools/highlight_code.py                       # ppt/*_portfolio.pptx 전부
    python ppt/tools/highlight_code.py ppt/wielder_portfolio.pptx
    python ppt/tools/highlight_code.py --dry-run             # 저장하지 않고 미리보기

코드 박스(Courier New 글꼴의 텍스트 상자)만 건드린다. 문단을 run 으로 쪼개 색만 입히고
글자·글꼴·크기·줄간격은 그대로 둔다. 여러 번 돌려도 결과가 같다.
"""
import copy
import glob
import io
import os
import re
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.dml.color import RGBColor

# ---------------------------------------------------------------- 팔레트
KEYWORD = '569CD6'
CONTROL = 'C586C0'
TYPE = '4EC9B0'
METHOD = 'DCDCAA'
IDENT = '9CDCFE'
STRING = 'CE9178'
NUMBER = 'B5CEA8'
COMMENT = '6A9955'
PLAIN = 'DCDCDC'

CONTROL_WORDS = {
    'if', 'else', 'elif', 'for', 'foreach', 'while', 'do', 'switch', 'case',
    'return', 'break', 'continue', 'yield', 'goto', 'throw',
}

KEYWORDS = {
    # C#
    'abstract', 'as', 'async', 'await', 'base', 'bool', 'byte', 'catch', 'char',
    'class', 'const', 'decimal', 'default', 'delegate', 'double', 'enum', 'event',
    'explicit', 'extern', 'false', 'finally', 'fixed', 'float', 'get', 'implicit',
    'in', 'int', 'interface', 'internal', 'is', 'lock', 'long', 'namespace', 'new',
    'null', 'object', 'operator', 'out', 'override', 'params', 'partial', 'private',
    'protected', 'public', 'readonly', 'ref', 'sbyte', 'sealed', 'set', 'short',
    'sizeof', 'static', 'string', 'struct', 'this', 'true', 'try', 'typeof', 'uint',
    'ulong', 'unsafe', 'ushort', 'using', 'value', 'virtual', 'void', 'volatile',
    'where',
    # Python
    'def', 'import', 'from', 'not', 'and', 'or', 'None', 'True', 'False', 'with',
    'except', 'raise', 'lambda', 'pass', 'global', 'self',
}

TOKEN_RE = re.compile(r"""
    (?P<comment>   //.*$ | \#.*$ )
  | (?P<string>    " (?: \\. | [^"\\] )* "? | ' (?: \\. | [^'\\] )* '? )
  | (?P<number>    \b 0[xX][0-9a-fA-F]+ | \b \d+ (?:\.\d+)? (?:[eE][+-]?\d+)? [fdmFDM]? \b )
  | (?P<ident>     [A-Za-z_][A-Za-z0-9_]* )
  | (?P<space>     \s+ )
  | (?P<other>     . )
""", re.VERBOSE)


def classify_block(lines):
    """여러 줄을 한꺼번에 처리한다. 파이썬 독스트링처럼 줄을 넘어가는 토큰 때문에 필요하다."""
    out = []
    in_doc = False
    for line in lines:
        fences = line.count('"""')
        if in_doc:
            out.append([(line, STRING)])
            if fences:
                in_doc = False
            continue
        if fences == 1:                      # 독스트링 시작
            out.append([(line, STRING)])
            in_doc = True
            continue
        if fences >= 2:                      # 한 줄짜리 독스트링
            out.append([(line, STRING)])
            continue
        out.append(classify(line))
    return out


def classify(line):
    """한 줄을 (텍스트, 색) 조각들로 쪼갠다. 조각을 이어 붙이면 원본과 같다."""
    out = []
    pos = 0
    prev_ident = None
    for m in TOKEN_RE.finditer(line):
        kind = m.lastgroup
        text = m.group()
        if kind == 'comment':
            out.append((text, COMMENT))
        elif kind == 'string':
            out.append((text, STRING))
        elif kind == 'number':
            out.append((text, NUMBER))
        elif kind == 'ident':
            out.append((text, None))          # 색은 뒤에서 결정 (다음 글자를 봐야 한다)
        elif kind == 'space':
            out.append((text, PLAIN))
        else:
            out.append((text, PLAIN))
        pos = m.end()

    # 식별자 색 결정 — 바로 뒤 문자와 바로 앞 문자를 본다.
    resolved = []
    for i, (text, color) in enumerate(out):
        if color is not None:
            resolved.append((text, color))
            continue
        nxt = ''
        for t2, _ in out[i + 1:]:
            if t2.strip():
                nxt = t2.strip()[0]
                break
        prev = ''
        prev_j = None
        for j in range(i - 1, -1, -1):
            if out[j][0].strip():
                prev = out[j][0].strip()[-1]
                prev_j = j
                break
        # 어트리뷰트는 '[' 가 줄의 첫 글자일 때만. attackers[i] 의 i 를 타입으로 보면 안 된다.
        is_attribute = False
        if prev == '[' and prev_j is not None:
            before = ''.join(t for t, _ in out[:prev_j + 1])
            is_attribute = before.strip() == '['
        if text in CONTROL_WORDS:
            resolved.append((text, CONTROL))
        elif text in KEYWORDS:
            resolved.append((text, KEYWORD))
        elif is_attribute:
            resolved.append((text, TYPE))          # 어트리뷰트
        elif nxt == '(':
            resolved.append((text, METHOD))
        elif prev == '.':
            resolved.append((text, IDENT))         # 멤버 접근 — cells.Length 는 속성이다
        elif text[0].isupper():
            resolved.append((text, TYPE))
        else:
            resolved.append((text, IDENT))

    # 같은 색끼리 합쳐 run 수를 줄인다.
    merged = []
    for text, color in resolved:
        if merged and merged[-1][1] == color:
            merged[-1] = (merged[-1][0] + text, color)
        else:
            merged.append((text, color))
    return merged


def is_code_box(shape):
    if not shape.has_text_frame:
        return False
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.name and 'Courier' in run.font.name:
                return True
    return False


def paint(shape):
    """문단마다 run 을 다시 만들어 색을 입힌다. 서식은 첫 run 것을 물려받는다."""
    paras = list(shape.text_frame.paragraphs)
    texts = [''.join(r.text for r in p.runs) for p in paras]
    blocks = classify_block(texts)
    changed = 0
    for para, text, pieces in zip(paras, texts, blocks):
        runs = para.runs
        if not runs:
            continue
        if not text.strip():
            # 빈 줄. 글자는 없지만 옛 색이 남지 않도록 기본색으로 맞춰 둔다.
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            para.runs[0].font.color.rgb = RGBColor.from_string(PLAIN)
            continue

        template = copy.deepcopy(runs[0]._r)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        keep = para.runs[0]
        keep.text = pieces[0][0]
        keep.font.color.rgb = RGBColor.from_string(pieces[0][1])
        anchor = keep._r
        for t, c in pieces[1:]:
            new_r = copy.deepcopy(template)
            anchor.addnext(new_r)
            anchor = new_r
            para.runs[-1]  # touch to refresh
            run = [r for r in para.runs if r._r is new_r][0]
            run.text = t
            run.font.color.rgb = RGBColor.from_string(c)
        changed += 1
    return changed


def process(path, dry_run=False):
    prs = Presentation(path)
    boxes = 0
    lines = 0
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not is_code_box(shape):
                continue
            boxes += 1
            n = paint(shape)
            lines += n
            print(f'  s{i} {shape.name}: {n} lines')
    if not dry_run and boxes:
        prs.save(path)
    print(f'{os.path.basename(path)}: {boxes} code boxes, {lines} lines'
          + ('  (dry-run)' if dry_run else '  saved'))
    return boxes


def main(argv):
    dry = '--dry-run' in argv
    args = [a for a in argv if not a.startswith('--')]
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    targets = args or sorted(glob.glob(os.path.join(here, '*_portfolio.pptx')))
    if not targets:
        print('대상 pptx 를 찾지 못했다.')
        return 1
    for t in targets:
        process(t, dry)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
